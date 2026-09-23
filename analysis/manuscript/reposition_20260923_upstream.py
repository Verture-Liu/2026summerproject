"""Reposition PaleoRigor as a pre-analytical control layer upstream of ancient
metagenomic profiling pipelines, cite those pipelines, and promote the
domain-specific refusal rules from Supplementary File 4 into the main text.

Adds three verified references after reference 18 and shifts references 19-34
to 22-37, remapping every in-text citation marker accordingly.
"""
from __future__ import annotations

import re
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "latest version_PaleoRigor_restructured.docx"
OUTPUT = ROOT / "latest version_PaleoRigor_upstream_positioned.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

SHIFT_FROM, SHIFT_BY = 19, 3

NEW_REFS = [
    "19. Pochon Z, Bergfeldt N, Kırdök E, Vicente M, Naidoo T, van der Valk T, et al. aMeta: an accurate and "
    "memory-efficient ancient metagenomic profiling workflow. Genome Biol. 2023;24:242. "
    "doi:10.1186/s13059-023-03083-9.",
    "20. Fellows Yates JA, Lamnidis TC, Borry M, Andrades Valtueña A, Fagernäs Z, Clayton S, et al. "
    "Reproducible, portable, and efficient ancient genome reconstruction with nf-core/eager. PeerJ. 2021;9:e10947. "
    "doi:10.7717/peerj.10947.",
    "21. Dhibar A, Matz MV. MAT-classifier: a memory-efficient pipeline for accurate genus level profiling from "
    "ancient metagenomic data. bioRxiv. 2026. doi:10.64898/2026.01.28.702372.",
]


def ptext(p):
    return "".join(n.text or "" for n in p.findall(f".//{W}t"))


def set_text(p, new):
    nodes = p.findall(f".//{W}t")
    old = "".join(n.text or "" for n in nodes)
    if old == new:
        return
    for tag, a, b, c, d in reversed(SequenceMatcher(None, old, new, autojunk=False).get_opcodes()):
        if tag == "equal":
            continue
        spans, pos = [], 0
        for n in nodes:
            v = n.text or ""
            spans.append((n, pos, pos + len(v)))
            pos += len(v)
        start = next((i for i, (_, x, y) in enumerate(spans) if x <= a < y), len(spans) - 1)
        if a == b:
            n, x, _ = spans[start]
            v = n.text or ""
            n.text = v[: a - x] + new[c:d] + v[a - x :]
        else:
            for j, (n, x, y) in enumerate([(n, x, y) for n, x, y in spans if x < b and y > a]):
                v = n.text or ""
                n.text = v[: max(0, a - x)] + (new[c:d] if j == 0 else "") + v[min(len(v), b - x) :]
        for n in nodes:
            n.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    assert ptext(p) == new, f"rewrite mismatch\nwant {new[:80]!r}\ngot  {ptext(p)[:80]!r}"


CITE = re.compile(r"\[(\d+(?:\s*[,–-]\s*\d+)*)\]")


def shift_citation(match: re.Match) -> str:
    inner = match.group(1)
    def bump(m):
        n = int(m.group(0))
        return str(n + SHIFT_BY if n >= SHIFT_FROM else n)
    return "[" + re.sub(r"\d+", bump, inner) + "]"


TITLE = ("PaleoRigor: expert-gated, auditable data preparation upstream of ancient metagenomic profiling")

EDITS = [
    # ---- Abstract -------------------------------------------------------
    ("Background: Ancient microbial evidence informs research on past health, diet, and environments, but its "
     "scarcity and susceptibility to contamination demand careful analysis. Paleobiologists need accessible tools to "
     "check file selection and processing decisions while retaining expert oversight.",
     "Background: Ancient microbial evidence informs research on past health, diet, and environments, but its "
     "scarcity and susceptibility to contamination demand careful analysis. Established pipelines profile ancient "
     "metagenomes once an operator has chosen the input files and the analysis to run; they do not check that choice, "
     "record what was changed before they started, or limit what may be concluded from quality-control output. These "
     "pre-analytical decisions are where an error becomes irrecoverable."),
    ("Conclusions: PaleoRigor supports workflow review and traceable data handling before specialist interpretation, "
     "with an Apple Silicon macOS application providing the tools for local use.",
     "Conclusions: PaleoRigor occupies the step before profiling. It produces verified inputs, a recorded "
     "transformation history, and an explicit statement of what the available evidence does not support, which an "
     "established ancient-metagenomic pipeline then consumes; an Apple Silicon macOS application provides the tools "
     "for local use."),
    # ---- Key points -----------------------------------------------------
    ("It makes these risks reviewable by linking each request to its inputs, approved operations, intermediate files, "
     "and expert decisions.",
     "It links each request to its inputs, approved operations, intermediate files and expert decisions, and encodes "
     "ancient-DNA-specific limits on interpretation — refusing, for example, any request to establish "
     "authenticity or absence of contamination from quality-control output."),
    # ---- Background: cite the domain pipelines and name the gap ----------
    ("Standardized workflows can reduce manual handling and make complex analyses easier to review. Workflow engines "
     "support reproducibility once users have specified inputs and operations [16–18]. Determining whether those "
     "choices suit the research question still requires expertise. Natural-language agents offer a way to connect "
     "requests with analysis tools [19], but their plans can contain nonexistent files, incompatible operations, or "
     "unsupported interpretations. What is missing is a mechanism that states a proposed analysis explicitly and "
     "constrains what it may do before it runs, while leaving responsibility for the biological conclusions with the "
     "researcher.",
     "Standardized workflows can reduce manual handling and make complex analyses easier to review. Generic workflow "
     "engines support reproducibility once users have specified inputs and operations [16–18], and the "
     "ancient-metagenomics community has built domain pipelines on them: aMeta implements profiling and "
     "authentication as a Snakemake workflow [19], nf-core/eager implements genome reconstruction in Nextflow [20], "
     "and more recent workflows continue to optimise the profiling step itself [21]. Each begins once an operator has "
     "decided which files to analyse and which analysis suits the question. Neither decision is checked by the "
     "pipeline, and neither is recorded by it. Natural-language agents offer a way to connect requests with analysis "
     "tools [22], but their plans can contain nonexistent files, incompatible operations, or unsupported "
     "interpretations. What is missing is a layer that runs before these pipelines: one that states a proposed "
     "analysis explicitly, constrains what it may do, records what it changed, and leaves responsibility for the "
     "biological conclusions with the researcher."),
    ("This division allows paleobiologists to review the proposed analysis and assess its results without delegating "
     "scientific judgment to the planning model (Figure 1; Tables 1 and 2).",
     "This division allows paleobiologists to review the proposed analysis and assess its results without delegating "
     "scientific judgment to the planning model (Figure 1; Tables 1 and 2). PaleoRigor is not a profiling pipeline "
     "and does not replace one: its outputs are the verified inputs, the recorded transformations and the audit "
     "trail that a pipeline such as aMeta or nf-core/eager then consumes [19, 20]."),
    # ---- Discussion: name the downstream neighbours ----------------------
    ("Source and workflow review can precede adapter removal, damage analysis, host removal, taxonomic profiling, "
     "contamination assessment, and statistical interpretation [13, 15].",
     "Source and workflow review can precede adapter removal, damage analysis, host removal, taxonomic profiling, "
     "contamination assessment, and statistical interpretation [13, 15], whether those steps are run by hand or "
     "through an established pipeline [19, 20]."),
]

# New paragraph promoting the domain-specific refusal rules into the main text,
# inserted after the expert-checkpoint paragraph in the first Results section.
ANCHOR_FOR_RULES = ("Expert review is placed at three points in that sequence: when the request is framed, before "
                    "execution, and when the retained evidence is interpreted.")
RULES_PARAGRAPH = (
    "Four of these limits are specific to ancient biomolecular work, and are stated to the planner as refusal rules "
    "rather than left to the user to remember. Quality-control output cannot establish that reads are ancient or free "
    "of contamination, so a request to demonstrate either is refused rather than answered. Damage analysis requires "
    "aligned data and an explicitly named reference, so it is refused when those prerequisites are absent. Read "
    "alignment and host removal are refused when no reference genome or existing index is named, instead of "
    "defaulting to one. A request whose file format contradicts the operation is refused rather than silently "
    "replaced by a different analysis. Each refusal returns a named reason code, and the rules are reproduced "
    "verbatim in Supplementary File 4."
)


def main():
    doc = Document(SOURCE)
    is_ref = lambda t: re.match(r"^\d+\.\s", t.strip()) is not None

    # Phase 1 - shift every in-text citation marker (body text and tables).
    shifted = 0
    for p in doc.paragraphs:
        t = ptext(p._p)
        if is_ref(t) or "[" not in t:
            continue
        new = CITE.sub(shift_citation, t)
        if new != t:
            set_text(p._p, new)
            shifted += 1
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    t = ptext(cp._p)
                    if "[" not in t:
                        continue
                    new = CITE.sub(shift_citation, t)
                    if new != t:
                        set_text(cp._p, new)
                        shifted += 1

    # Phase 2 - renumber the reference list, high to low.
    refs = [p for p in doc.paragraphs if is_ref(p.text)]
    assert len(refs) == 34, f"expected 34 references, found {len(refs)}"
    for p in reversed(refs):
        t = ptext(p._p)
        n = int(re.match(r"^(\d+)\.", t.strip()).group(1))
        if n >= SHIFT_FROM:
            set_text(p._p, re.sub(r"^(\s*)\d+\.", lambda m: f"{m.group(1)}{n + SHIFT_BY}.", t, count=1))

    # Phase 3 - insert the three new references after reference 18.
    ref18 = next(p for p in doc.paragraphs if ptext(p._p).strip().startswith("18. Sandve"))
    anchor = ref18
    for text in NEW_REFS:
        el = OxmlElement("w:p")
        anchor._p.addnext(el)
        para = Paragraph(el, anchor._parent)
        para.style = ref18.style
        para.add_run(text)
        anchor = para

    # Phase 4 - positioning edits.
    for old, new in EDITS:
        # The document's citations were already shifted in phase 1, so shift the
        # search text the same way. `new` is written with final numbering.
        old_shifted = CITE.sub(shift_citation, old)
        hits = [p for p in doc.paragraphs if old_shifted in ptext(p._p)]
        assert len(hits) == 1, f"expected one match, found {len(hits)} for {old_shifted[:60]!r}"
        p = hits[0]._p
        set_text(p, ptext(p).replace(old_shifted, new))

    # Phase 5 - title and the promoted refusal rules.
    title = doc.paragraphs[0]
    assert "PaleoRigor:" in title.text, "first paragraph is not the title"
    set_text(title._p, TITLE)

    anchor = next(p for p in doc.paragraphs if ptext(p._p).startswith(ANCHOR_FOR_RULES))
    el = OxmlElement("w:p")
    anchor._p.addnext(el)
    para = Paragraph(el, anchor._parent)
    para.style = anchor.style
    para.add_run(RULES_PARAGRAPH)

    doc.core_properties.title = TITLE
    doc.save(OUTPUT)

    # ----------------------------------------------------------- verification
    out = Document(OUTPUT)
    checks = []
    refs_out = [p.text.strip() for p in out.paragraphs if is_ref(p.text)]
    nums = [int(re.match(r"^(\d+)\.", r).group(1)) for r in refs_out]
    assert nums == list(range(1, 38)), f"reference numbering broken: {nums}"
    checks.append(f"{len(refs_out)} references numbered 1-37 without gaps")

    for key, ref_no in [("aMeta", 19), ("nf-core/eager", 20), ("MAT-classifier", 21)]:
        assert refs_out[ref_no - 1].startswith(f"{ref_no}."), f"reference {ref_no} misplaced"
        assert key in refs_out[ref_no - 1], f"reference {ref_no} is not {key}"
    checks.append("aMeta = 19, nf-core/eager = 20, MAT-classifier = 21")

    body = "\n".join(p.text for p in out.paragraphs if not is_ref(p.text))
    tbls = "\n".join(c.text for t in out.tables for r in t.rows for c in r.cells)
    cited = set()
    for m in CITE.finditer(body + "\n" + tbls):
        for part in m.group(1).split(","):
            part = part.strip()
            span = re.match(r"^(\d+)\s*[\u2013-]\s*(\d+)$", part)
            if span:
                cited.update(range(int(span.group(1)), int(span.group(2)) + 1))
            elif part.isdigit():
                cited.add(int(part))
    missing = sorted(set(range(1, 38)) - cited)
    assert not missing, f"references never cited: {missing}"
    checks.append("every reference 1-37 is cited in the text")
    assert max(cited) <= 37, f"citation exceeds reference list: {max(cited)}"
    checks.append(f"no citation exceeds 37 (highest cited: {max(cited)})")

    assert out.paragraphs[0].text == TITLE
    checks.append("title updated")
    assert "refused rather than answered" in body
    checks.append("domain-specific refusal rules promoted into the main text")
    assert "aMeta implements profiling and authentication as a Snakemake workflow" in body
    checks.append("Background names the downstream pipelines and the gap they leave")

    print("\n".join(f"  OK  {c}" for c in checks))
    print(f"\n  citation markers shifted in {shifted} paragraphs/cells")
    print(f"\nWrote {OUTPUT.name}")


if __name__ == "__main__":
    main()
