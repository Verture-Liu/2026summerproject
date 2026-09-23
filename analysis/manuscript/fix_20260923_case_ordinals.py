"""Remove case ordinals from the main text and explain the supplementary
case numbering, which follows the repository layout rather than the order of
presentation."""
from __future__ import annotations

from difflib import SequenceMatcher
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "latest version_PaleoRigor_restructured.docx"
OUTPUT = ROOT / "latest version_PaleoRigor_restructured.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def ptext(p):
    return "".join(n.text or "" for n in p.findall(f".//{W}t"))


def set_text(p, new):
    nodes = p.findall(f".//{W}t")
    old = "".join(n.text or "" for n in nodes)
    for tag, a, b, c, d in reversed(SequenceMatcher(None, old, new, autojunk=False).get_opcodes()):
        if tag == "equal":
            continue
        spans, pos = [], 0
        for n in nodes:
            v = n.text or ""
            spans.append((n, pos, pos + len(v)))
            pos += len(v)
        start = next((i for i, (_, x, y) in enumerate(spans) if x <= a < y), len(spans) - 1)
        if a == b:
            n, x, _ = spans[start]
            v = n.text or ""
            n.text = v[: a - x] + new[c:d] + v[a - x :]
        else:
            for j, (n, x, y) in enumerate([(n, x, y) for n, x, y in spans if x < b and y > a]):
                v = n.text or ""
                n.text = v[: max(0, a - x)] + (new[c:d] if j == 0 else "") + v[min(len(v), b - x) :]
        for n in nodes:
            n.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    assert ptext(p) == new


EDITS = [
    ("The second case tested whether local sequencing files agreed",
     "This case tested whether local sequencing files agreed"),
    ("The first case tested whether saved records could explain",
     "This case tested whether saved records could explain"),
    ("The third case examined how expert review could limit conclusions",
     "This case examined how expert review could limit conclusions"),
    ("Figure 3 examines source agreement, Figure 4 traces a recorded transformation, and Figure 5 shows where "
     "file-level evidence stops.",
     "Figure 3 examines source agreement, Figure 4 traces a recorded transformation, and Figure 5 shows where "
     "file-level evidence stops. Supplementary Tables S1–S3 label these cases by their repository folders "
     "— Case 1, the peptide table; Case 2, the sequencing records; Case 3, the retraction-associated audit "
     "— following the order in which they were run rather than the order presented here."),
]


def main():
    doc = Document(SOURCE)
    for old, new in EDITS:
        hits = [p for p in doc.paragraphs if old in ptext(p._p)]
        assert len(hits) == 1, f"expected one match for {old[:50]!r}, found {len(hits)}"
        p = hits[0]._p
        set_text(p, ptext(p).replace(old, new))
    doc.save(OUTPUT)

    out = Document(OUTPUT)
    text = "\n".join(p.text for p in out.paragraphs)
    for stale in ["The first case", "The second case", "The third case"]:
        assert stale not in text, f"ordinal remains: {stale}"
    assert "following the order in which they were run" in text
    print("  OK  case ordinals removed from the main text")
    print("  OK  supplementary case numbering explained")
    print(f"\nUpdated {OUTPUT.name}")


if __name__ == "__main__":
    main()
