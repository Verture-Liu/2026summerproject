"""Insert the minted Zenodo DOI, release tag and commit into the manuscript.

Values come from revision_20260923_evidence.json, whose DOI entries record the
Zenodo record they were read from. The script refuses to run if any placeholder
it expects is absent, so it cannot silently write a DOI into the wrong place.
"""
from __future__ import annotations

import json
import re
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "latest version_PaleoRigor_ablation_disclosure.docx"
OUTPUT = ROOT / "latest version_PaleoRigor_archived_release.docx"
E = {k: v["value"] for k, v in
     json.loads((ROOT / "analysis/manuscript/revision_20260923_evidence.json").read_text(encoding="utf-8")).items()}

TAG, COMMIT = E["release_tag"], E["release_commit"]
VDOI, CDOI = E["zenodo_version_doi"], E["zenodo_concept_doi"]


def ptext(p) -> str:
    return "".join(n.text or "" for n in p.xpath(".//w:t"))


def replace_preserving_runs(p, new: str) -> None:
    nodes = p.xpath(".//w:t")
    old = "".join(n.text or "" for n in nodes)
    for tag, a, b, c, d in reversed(SequenceMatcher(None, old, new, autojunk=False).get_opcodes()):
        if tag == "equal":
            continue
        spans, pos = [], 0
        for n in nodes:
            val = n.text or ""
            spans.append((n, pos, pos + len(val)))
            pos += len(val)
        start = next((i for i, (_, x, y) in enumerate(spans) if x <= a < y), len(spans) - 1)
        if a == b:
            n, x, _ = spans[start]
            val = n.text or ""
            n.text = val[: a - x] + new[c:d] + val[a - x :]
        else:
            touched = [(i, n, x, y) for i, (n, x, y) in enumerate(spans) if x < b and y > a]
            for j, (i, n, x, y) in enumerate(touched):
                val = n.text or ""
                n.text = val[: max(0, a - x)] + (new[c:d] if j == 0 else "") + val[min(len(val), b - x) :]
        for n in nodes:
            n.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    assert ptext(p) == new


EDITS = [
    ("Archived version: [Zenodo DOI and release tag to be inserted before submission].",
     f"Archived version: release tag {TAG} (commit {COMMIT}), archived at Zenodo under DOI {VDOI}; "
     f"the concept DOI {CDOI} always resolves to the latest version."),
    ("Repository: PaleoRigor GitHub repository [31]; versioned release tag and commit to be assigned before submission.",
     f"Repository: PaleoRigor GitHub repository [31]; release tag {TAG}, commit {COMMIT}."),
    ("To be archived as a tagged GitHub release and DOI-linked snapshot before submission.",
     f"Archived as a tagged GitHub release with Zenodo DOI {VDOI}."),
    ("GitHub repository: PaleoRigor GitHub repository [31]; planned archive: [Zenodo DOI pending]",
     f"GitHub repository: PaleoRigor GitHub repository [31]; archived release: Zenodo DOI {VDOI}"),
    ("Replace mutable main-branch links with release-specific links before final submission.",
     f"Reported results correspond to release {TAG}; cite the archived snapshot rather than the mutable main branch."),
]


def main() -> None:
    doc = Document(SOURCE)
    targets = list(doc.paragraphs) + [
        p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs
    ]
    applied = []
    for old, new in EDITS:
        hits = [p for p in targets if old in ptext(p._p)]
        assert len(hits) == 1, f"expected exactly one paragraph containing {old[:60]!r}, found {len(hits)}"
        p = hits[0]._p
        before = ptext(p)
        replace_preserving_runs(p, before.replace(old, new))
        applied.append((before, ptext(p)))

    core = doc.core_properties
    core.subject = (f"Archived release {TAG} (Zenodo DOI {VDOI}). " + (core.subject or ""))[:255]
    doc.save(OUTPUT)

    out = Document(OUTPUT)
    text = "\n".join(p.text for p in out.paragraphs) + "\n".join(
        c.text for t in out.tables for r in t.rows for c in r.cells)
    leftovers = [s for s in ("Zenodo DOI pending", "to be assigned before submission",
                             "to be inserted before submission", "planned archive") if s in text]
    assert not leftovers, f"placeholders remain: {leftovers}"
    assert text.count(VDOI) >= 3, "version DOI not written in every expected place"
    assert CDOI in text, "concept DOI missing"
    src_refs = [p.text for p in Document(SOURCE).paragraphs if re.match(r"^\d+\.\s", p.text.strip())]
    out_refs = [p.text for p in out.paragraphs if re.match(r"^\d+\.\s", p.text.strip())]
    assert src_refs == out_refs, "reference list changed"

    print(f"  OK  {len(applied)} placeholders replaced")
    print(f"  OK  version DOI {VDOI} present {text.count(VDOI)}x, concept DOI present")
    print(f"  OK  no archive placeholders remain; {len(out_refs)} references unchanged")
    for before, after in applied:
        print(f"\n- {before[:150]}\n+ {after[:230]}")
    print(f"\nWrote {OUTPUT.name}")


if __name__ == "__main__":
    main()
