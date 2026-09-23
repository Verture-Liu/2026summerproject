"""Add four verified ancient-DNA references that the manuscript needs, and
renumber the reference list and every in-text citation accordingly.

Each reference was checked against its publisher or index record before being
written here; none is inferred. DamageProfiler in particular is named in the
manuscript's own refusal rules and was previously uncited.
"""
from __future__ import annotations

import re
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "latest version_PaleoRigor_upstream_positioned.docx"
OUTPUT = ROOT / "latest version_PaleoRigor_paleo_refs.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# old number -> new number. Monotonic, so the existing list order is preserved.
REMAP = {}
for old in range(1, 8):
    REMAP[old] = old                 # 1-7 unchanged
for old in range(8, 13):
    REMAP[old] = old + 1             # 8-12  -> 9-13   (AncientMetagenomeDir becomes 8)
for old in range(13, 16):
    REMAP[old] = old + 2             # 13-15 -> 15-17  (Key 2017 becomes 14)
for old in range(16, 21):
    REMAP[old] = old + 3             # 16-20 -> 19-23  (DamageProfiler becomes 18)
for old in range(21, 38):
    REMAP[old] = old + 4             # 21-37 -> 25-41  (HOPS becomes 24)

# text, and the old reference number it is inserted after
NEW_REFS = [
    (7, "8. Fellows Yates JA, Andrades Valtueña A, Vågene ÅJ, Cribdon B, Velsko IM, Borry M, et al. "
        "Community-curated and standardised metadata of published ancient metagenomic samples with "
        "AncientMetagenomeDir. Sci Data. 2021;8:31. doi:10.1038/s41597-021-00816-y."),
    (12, "14. Key FM, Posth C, Krause J, Herbig A, Bos KI. Mining metagenomic data sets for ancient DNA: recommended "
         "protocols for authentication. Trends Genet. 2017;33:508–520. doi:10.1016/j.tig.2017.05.005."),
    (15, "18. Neukamm J, Peltzer A, Nieselt K. DamageProfiler: fast damage pattern calculation for ancient DNA. "
         "Bioinformatics. 2021;37:3652–3653. doi:10.1093/bioinformatics/btab190."),
    (20, "24. Hübler R, Key FM, Warinner C, Bos KI, Krause J, Herbig A. HOPS: automated detection and "
         "authentication of pathogen DNA in archaeological remains. Genome Biol. 2019;20:280. "
         "doi:10.1186/s13059-019-1903-0."),
]


def ptext(p):
    return "".join(n.text or "" for n in p.findall(f".//{W}t"))


def set_text(p, new):
    nodes = p.findall(f".//{W}t")
    old = "".join(n.text or "" for n in nodes)
    if old == new:
        return
    for tag, a, b, c, d in reversed(SequenceMatcher(None, old, new, autojunk=False).get_opcodes()):
        if tag == "equal":
            continue
        spans, pos = [], 0
        for n in nodes:
            v = n.text or ""
            spans.append((n, pos, pos + len(v)))
            pos += len(v)
        start = next((i for i, (_, x, y) in enumerate(spans) if x <= a < y), len(spans) - 1)
        if a == b:
            n, x, _ = spans[start]
            v = n.text or ""
            n.text = v[: a - x] + new[c:d] + v[a - x :]
        else:
            for j, (n, x, y) in enumerate([(n, x, y) for n, x, y in spans if x < b and y > a]):
                v = n.text or ""
                n.text = v[: max(0, a - x)] + (new[c:d] if j == 0 else "") + v[min(len(v), b - x) :]
        for n in nodes:
            n.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    assert ptext(p) == new


CITE = re.compile(r"\[(\d+(?:\s*[,–-]\s*\d+)*)\]")


def remap_marker(m: re.Match) -> str:
    def one(x):
        n = int(x)
        assert n in REMAP, f"citation [{n}] has no mapping"
        return str(REMAP[n])
    parts = []
    for part in m.group(1).split(","):
        part = part.strip()
        span = re.match(r"^(\d+)\s*[–-]\s*(\d+)$", part)
        if span:
            parts.append(f"{one(span.group(1))}–{one(span.group(2))}")
        else:
            parts.append(one(part))
    return "[" + ", ".join(parts) + "]"


# Text edits that cite the new references. Written with FINAL numbering; the
# `old` side is remapped before matching, exactly as the document will have been.
EDITS = [
    ("Sequencing data are often deposited in the Sequence Read Archive (SRA) and European Nucleotide Archive (ENA) "
     "[6, 7].",
     "Sequencing data are often deposited in the Sequence Read Archive (SRA) and European Nucleotide Archive (ENA) "
     "[6, 7], and community curation has standardised the sample-level metadata that accompanies published ancient "
     "metagenomic records [8]."),
    ("Authenticity depends on several sources of evidence, including sample history, molecular damage, contamination "
     "controls, and archaeological context [3, 8–12].",
     "Authenticity depends on several sources of evidence, including sample history, molecular damage, contamination "
     "controls, and archaeological context [3, 9–13], and the field has published explicit protocols for "
     "authenticating ancient DNA recovered from metagenomic data [14]."),
    ("Mapping, adapter removal, and damage analysis assess the files supplied to them [13–15];",
     "Mapping, adapter removal, and damage analysis assess the files supplied to them [15–18];"),
    ("aMeta implements profiling and authentication as a Snakemake workflow [19], nf-core/eager implements genome "
     "reconstruction in Nextflow [20], and more recent workflows continue to optimise the profiling step itself [21].",
     "aMeta implements profiling and authentication as a Snakemake workflow [22], nf-core/eager implements genome "
     "reconstruction in Nextflow [23], HOPS couples pathogen screening to automated authentication [24], and more "
     "recent workflows continue to optimise the profiling step itself [25]."),
    ("Damage analysis requires aligned data and an explicitly named reference, so it is refused when those "
     "prerequisites are absent.",
     "Damage analysis requires aligned data and an explicitly named reference [15, 18], so it is refused when those "
     "prerequisites are absent."),
]


def main():
    doc = Document(SOURCE)
    is_ref = lambda t: re.match(r"^\d+\.\s", t.strip()) is not None

    refs = [p for p in doc.paragraphs if is_ref(p.text)]
    assert len(refs) == 37, f"expected 37 references, found {len(refs)}"
    old_numbers = [int(re.match(r"^(\d+)\.", p.text.strip()).group(1)) for p in refs]
    assert old_numbers == list(range(1, 38)), "reference list is not 1-37"

    # Phase 1 - remap every in-text citation marker.
    remapped = 0
    for p in doc.paragraphs:
        t = ptext(p._p)
        if is_ref(t) or "[" not in t:
            continue
        new = CITE.sub(remap_marker, t)
        if new != t:
            set_text(p._p, new)
            remapped += 1
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    t = ptext(cp._p)
                    if "[" not in t:
                        continue
                    new = CITE.sub(remap_marker, t)
                    if new != t:
                        set_text(cp._p, new)
                        remapped += 1

    # Phase 2 - renumber the existing reference entries.
    for p, old in zip(refs, old_numbers):
        t = ptext(p._p)
        set_text(p._p, re.sub(r"^(\s*)\d+\.", lambda m: f"{m.group(1)}{REMAP[old]}.", t, count=1))

    # Phase 3 - insert the new references after their anchors.
    by_old = dict(zip(old_numbers, refs))
    for after_old, text in NEW_REFS:
        anchor = by_old[after_old]
        el = OxmlElement("w:p")
        anchor._p.addnext(el)
        para = Paragraph(el, anchor._parent)
        para.style = anchor.style
        para.add_run(text)

    # Phase 4 - add the citations themselves.
    for old, new in EDITS:
        old_remapped = CITE.sub(remap_marker, old)
        hits = [p for p in doc.paragraphs if old_remapped in ptext(p._p)]
        assert len(hits) == 1, f"expected one match, found {len(hits)} for {old_remapped[:70]!r}"
        p = hits[0]._p
        set_text(p, ptext(p).replace(old_remapped, new))

    doc.save(OUTPUT)

    # ------------------------------------------------------------ verification
    out = Document(OUTPUT)
    checks = []
    refs_out = [p.text.strip() for p in out.paragraphs if is_ref(p.text)]
    nums = [int(re.match(r"^(\d+)\.", r).group(1)) for r in refs_out]
    assert nums == list(range(1, 42)), f"numbering broken: {nums}"
    checks.append(f"{len(refs_out)} references numbered 1-41 without gaps")

    for n, key in [(8, "AncientMetagenomeDir"), (14, "Mining metagenomic data sets"),
                   (18, "DamageProfiler"), (24, "HOPS")]:
        assert key in refs_out[n - 1], f"reference {n} is not {key}: {refs_out[n-1][:70]}"
    checks.append("AncientMetagenomeDir = 8, Key 2017 = 14, DamageProfiler = 18, HOPS = 24")

    body = "\n".join(p.text for p in out.paragraphs if not is_ref(p.text))
    tbls = "\n".join(c.text for t in out.tables for r in t.rows for c in r.cells)
    cited = set()
    for m in CITE.finditer(body + "\n" + tbls):
        for part in m.group(1).split(","):
            part = part.strip()
            span = re.match(r"^(\d+)\s*[–-]\s*(\d+)$", part)
            if span:
                cited.update(range(int(span.group(1)), int(span.group(2)) + 1))
            elif part.isdigit():
                cited.add(int(part))
    missing = sorted(set(range(1, 42)) - cited)
    assert not missing, f"references never cited: {missing}"
    assert max(cited) <= 41, f"citation exceeds list: {max(cited)}"
    checks.append(f"every reference 1-41 is cited; highest citation is {max(cited)}")

    assert "DamageProfiler" in "\n".join(refs_out), "DamageProfiler still uncited"
    checks.append("DamageProfiler, named in the refusal rules, is now cited")

    print("\n".join(f"  OK  {c}" for c in checks))
    print(f"\n  citation markers remapped in {remapped} paragraphs/cells")
    print(f"\nWrote {OUTPUT.name}")


if __name__ == "__main__":
    main()
