"""Build the conservative variant: identical structure, identical references and
identical refusal-rule text as the repositioned version, but with the original
title and framing. The two files differ only in how the work is positioned, so
the supervisor can judge that one decision in isolation.
"""
from __future__ import annotations

from difflib import SequenceMatcher
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "latest version_PaleoRigor_paleo_refs.docx"
OUTPUT = ROOT / "VARIANT-A_original-framing.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def ptext(p):
    return "".join(n.text or "" for n in p.findall(f".//{W}t"))


def set_text(p, new):
    nodes = p.findall(f".//{W}t")
    old = "".join(n.text or "" for n in nodes)
    if old == new:
        return
    for tag, a, b, c, d in reversed(SequenceMatcher(None, old, new, autojunk=False).get_opcodes()):
        if tag == "equal":
            continue
        spans, pos = [], 0
        for n in nodes:
            v = n.text or ""
            spans.append((n, pos, pos + len(v)))
            pos += len(v)
        start = next((i for i, (_, x, y) in enumerate(spans) if x <= a < y), len(spans) - 1)
        if a == b:
            n, x, _ = spans[start]
            v = n.text or ""
            n.text = v[: a - x] + new[c:d] + v[a - x :]
        else:
            for j, (n, x, y) in enumerate([(n, x, y) for n, x, y in spans if x < b and y > a]):
                v = n.text or ""
                n.text = v[: max(0, a - x)] + (new[c:d] if j == 0 else "") + v[min(len(v), b - x) :]
        for n in nodes:
            n.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    assert ptext(p) == new


TITLE_A = "PaleoRigor: expert-guided workflows for traceable paleomicrobiome data processing"

REVERTS = [
    # Abstract Background
    ("Background: Ancient microbial evidence informs research on past health, diet, and environments, but its "
     "scarcity and susceptibility to contamination demand careful analysis. Established pipelines profile ancient "
     "metagenomes once an operator has chosen the input files and the analysis to run; they do not check that choice, "
     "record what was changed before they started, or limit what may be concluded from quality-control output. These "
     "pre-analytical decisions are where an error becomes irrecoverable.",
     "Background: Ancient microbial evidence informs research on past health, diet, and environments, but its "
     "scarcity and susceptibility to contamination demand careful analysis. Paleobiologists need accessible tools to "
     "check file selection and processing decisions while retaining expert oversight."),
    # Abstract Conclusions
    ("Conclusions: PaleoRigor occupies the step before profiling. It produces verified inputs, a recorded "
     "transformation history, and an explicit statement of what the available evidence does not support, which an "
     "established ancient-metagenomic pipeline then consumes; an Apple Silicon macOS application provides the tools "
     "for local use.",
     "Conclusions: PaleoRigor supports workflow review and traceable data handling before specialist interpretation, "
     "with an Apple Silicon macOS application providing the tools for local use."),
    # Background: keep the citations, drop the upstream framing
    ("Each begins once an operator has decided which files to analyse and which analysis suits the question. Neither "
     "decision is checked by the pipeline, and neither is recorded by it. Natural-language agents offer a way to "
     "connect requests with analysis tools [26], but their plans can contain nonexistent files, incompatible "
     "operations, or unsupported interpretations. What is missing is a layer that runs before these pipelines: one "
     "that states a proposed analysis explicitly, constrains what it may do, records what it changed, and leaves "
     "responsibility for the biological conclusions with the researcher.",
     "Determining whether those choices suit the research question still requires expertise. Natural-language agents "
     "offer a way to connect requests with analysis tools [26], but their plans can contain nonexistent files, "
     "incompatible operations, or unsupported interpretations. What is missing is a mechanism that states a proposed "
     "analysis explicitly and constrains what it may do before it runs, while leaving responsibility for the "
     "biological conclusions with the researcher."),
    # Background: drop the explicit "not a pipeline / feeds a pipeline" sentence
    (" PaleoRigor is not a profiling pipeline and does not replace one: its outputs are the verified inputs, the "
     "recorded transformations and the audit trail that a pipeline such as aMeta or nf-core/eager then consumes "
     "[22, 23].", ""),
    # Discussion: drop the downstream-pipeline clause
    (", whether those steps are run by hand or through an established pipeline [22, 23].", "."),
]


def main():
    doc = Document(SOURCE)
    title = doc.paragraphs[0]
    assert "PaleoRigor:" in title.text
    set_text(title._p, TITLE_A)
    doc.core_properties.title = TITLE_A

    for old, new in REVERTS:
        hits = [p for p in doc.paragraphs if old in ptext(p._p)]
        assert len(hits) == 1, f"expected one match, found {len(hits)} for {old[:60]!r}"
        p = hits[0]._p
        set_text(p, ptext(p).replace(old, new))

    doc.save(OUTPUT)

    a, b = Document(OUTPUT), Document(SOURCE)
    ta = [p.text for p in a.paragraphs]
    tb = [p.text for p in b.paragraphs]
    assert len(ta) == len(tb), "paragraph counts differ between variants"
    diffs = [i for i, (x, y) in enumerate(zip(ta, tb)) if x != y]
    refs_a = [p.text for p in a.paragraphs if p.text.strip()[:3].split(".")[0].isdigit()
              and p.text.strip().split(".")[0].isdigit()]
    print(f"  OK  both variants have {len(ta)} paragraphs")
    print(f"  OK  they differ in exactly {len(diffs)} paragraphs: {diffs}")
    for i in diffs:
        print(f"\n--- paragraph {i} ---")
        print(f"A (original framing): {ta[i][:200]}")
        print(f"B (upstream framing): {tb[i][:200]}")
    print(f"\nWrote {OUTPUT.name}")


if __name__ == "__main__":
    main()
